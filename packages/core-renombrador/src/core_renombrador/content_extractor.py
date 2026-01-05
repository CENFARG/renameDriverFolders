"""
ContentExtractor con soporte para OCR
======================================

Extrae contenido de archivos de diversos formatos, incluyendo:
- Texto plano
- PDF (con texto y escaneados)
- DOCX
- XLSX
- Imágenes (JPG, PNG) mediante OCR

:created:   2025-12-05
:filename:  content_extractor.py
:author:    amBotHs + CENF
:version:   2.0.0
:status:    Development
:license:   MIT
:copyright: Copyright (c) 2025 CENF
"""

import logging
import os
from io import BytesIO
from typing import Optional, Tuple

import docx
import openpyxl
from google.cloud import vision
from pdf2image import convert_from_bytes
from PIL import Image
from pypdf import PdfReader

logger = logging.getLogger(__name__)


class ContentExtractor:
    """
    Extrae texto de archivos de múltiples formatos.
    Supports OCR for images and scanned PDFs using Google Cloud Vision.
    """

    def __init__(self, enable_ocr: bool = True, min_text_threshold: int = 100):
        """
        Initialize ContentExtractor.

        Args:
            enable_ocr: Enable OCR for images and scanned PDFs.
                        Habilitar OCR para imágenes y PDFs escaneados.
            min_text_threshold: Minimum text length to consider PDF as "text-based".
                                Longitud mínima de texto para considerar un PDF como "con texto".
        """
        self.enable_ocr = enable_ocr
        self.min_text_threshold = min_text_threshold
        
        if self.enable_ocr:
            try:
                self.vision_client = vision.ImageAnnotatorClient()
                logger.info("Google Cloud Vision client initialized successfully")
            except Exception as e:
                logger.warning(f"Could not initialize Vision API client: {e}. OCR disabled.")
                self.vision_client = None
                self.enable_ocr = False

    def get_content(self, file_path: str, file_bytes: bytes) -> str:
        """
        Extracts text content from a file based on its extension.
        Extrae contenido de texto de un archivo según su extensión.

        Args:
            file_path: File name or path.
                      Nombre o ruta del archivo.
            file_bytes: File content as bytes.
                       Contenido del archivo como bytes.

        Returns:
            Extracted text content.
            Contenido de texto extraído.
        """
        _, extension = os.path.splitext(file_path)
        extension = extension.lower()

        try:
            if extension == ".txt":
                return file_bytes.decode("utf-8", errors="ignore")
            elif extension == ".xlsx":
                return self._get_xlsx_content(file_bytes)
            elif extension == ".docx":
                return self._get_docx_content(file_bytes)
            elif extension == ".pdf":
                return self._get_pdf_content(file_bytes)
            elif extension in [".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tiff"]:
                return self._get_image_content(file_bytes)
            else:
                # Fallback: try to decode as text
                try:
                    return file_bytes.decode("utf-8", errors="ignore")
                except Exception:
                    return "[Unsupported file type]"
        except Exception as e:
            logger.error(f"Error extracting content from {file_path}: {e}")
            return f"[Error extracting content: {str(e)}]"

    def _get_xlsx_content(self, file_bytes: bytes) -> str:
        """Extracts content from an XLSX file."""
        workbook = openpyxl.load_workbook(BytesIO(file_bytes))
        text = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value:
                        text.append(str(cell.value))
        return "\\n".join(text)

    def _get_docx_content(self, file_bytes: bytes) -> str:
        """Extracts content from a DOCX file."""
        doc = docx.Document(BytesIO(file_bytes))
        text = []
        for para in doc.paragraphs:
            text.append(para.text)
        return "\\n".join(text)

    def _get_pdf_content(self, file_bytes: bytes) -> str:
        """
        Extracts content from a PDF file.
        First tries text extraction. If insufficient text is found, uses OCR.
        """
        # Try text extraction first
        try:
            reader = PdfReader(BytesIO(file_bytes))
            text = []
            for page in reader.pages:
                page_text = page.extract_text() or ""
                text.append(page_text)
            
            combined_text = "\\n".join(text)
            
            # Check if we got enough text
            if len(combined_text.strip()) >= self.min_text_threshold:
                logger.debug(f"PDF text extraction successful ({len(combined_text)} chars)")
                return combined_text
            
            # If insufficient text and OCR is enabled, try OCR
            if self.enable_ocr and self.vision_client:
                logger.info("PDF appears to be scanned. Attempting OCR...")
                return self._ocr_pdf(file_bytes)
            else:
                logger.warning("Insufficient text extracted from PDF and OCR is disabled")
                return combined_text
                
        except Exception as e:
            logger.error(f"Error extracting PDF content: {e}")
            return "[Error extracting PDF content]"

    def _get_image_content(self, file_bytes: bytes) -> str:
        """
        Extracts text from image files using OCR.
        Extrae texto de archivos de imagen usando OCR.
        """
        if not self.enable_ocr or not self.vision_client:
            logger.warning("OCR is disabled. Cannot extract text from images.")
            return "[OCR disabled - image content not extracted]"
        
        try:
            return self._ocr_image_bytes(file_bytes)
        except Exception as e:
            logger.error(f"Error extracting image content: {e}")
            return f"[Error extracting image content: {str(e)}]"

    def _ocr_pdf(self, pdf_bytes: bytes) -> str:
        """
        Converts PDF pages to images and performs OCR.
        Convierte páginas PDF a imágenes y realiza OCR.
        """
        try:
            # Convert PDF to images
            images = convert_from_bytes(pdf_bytes, dpi=200, fmt="png")
            logger.info(f"Converted PDF to {len(images)} images for OCR")
            
            texts = []
            for i, image in enumerate(images):
                # Convert PIL Image to bytes
                img_byte_arr = BytesIO()
                image.save(img_byte_arr, format="PNG")
                img_bytes = img_byte_arr.getvalue()
                
                # Perform OCR
                page_text = self._ocr_image_bytes(img_bytes)
                texts.append(page_text)
                logger.debug(f"OCR completed for page {i+1}/{len(images)}")
            
            return "\\n".join(texts)
        
        except Exception as e:
            logger.error(f"Error performing OCR on PDF: {e}")
            return f"[Error performing OCR on PDF: {str(e)}]"

    def _ocr_image_bytes(self, image_bytes: bytes) -> str:
        """
        Performs OCR on image bytes using Google Cloud Vision.
        Realiza OCR en bytes de imagen usando Google Cloud Vision.
        """
        if not self.vision_client:
            raise ValueError("Vision API client not initialized")
        
        try:
            image = vision.Image(content=image_bytes)
            response = self.vision_client.document_text_detection(image=image)
            
            if response.error.message:
                raise Exception(f"Vision API error: {response.error.message}")
            
            # Get full text annotation
            if response.full_text_annotation:
                text = response.full_text_annotation.text
                logger.debug(f"OCR extracted {len(text)} characters")
                return text
            else:
                logger.warning("No text detected in image")
                return "[No text detected]"
                
        except Exception as e:
            logger.error(f"OCR failed: {e}")
            raise

    def get_content_with_confidence(
        self, file_path: str, file_bytes: bytes
    ) -> Tuple[str, Optional[float]]:
        """
        Extrae contenido y devuelve también un índice de confianza (para OCR).
        
        Returns:
            Tuple of (text, confidence_score)
            confidence_score is None for non-OCR extractions
        """
        # For now, this is a placeholder for future "verifiability index"
        # TODO: Implement confidence scoring based on OCR results
        text = self.get_content(file_path, file_bytes)
        return text, None  # Future: calculate confidence from Vision API
